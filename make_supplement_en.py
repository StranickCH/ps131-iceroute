# -*- coding: utf-8 -*-
"""Supplementary material (English): Fig. S1 from runs/drlv3_* and the docx with Tables S1-S7.
Table values are transcribed from the result files listed in README.md (no computation here).
Run: python make_supplement_en.py [runs folder] [output folder]"""
import os
import sys
import glob
import pandas as pd
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm

RUNS = sys.argv[1] if len(sys.argv) > 1 else "runs"
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(RUNS, "supplement_en")
os.makedirs(OUT, exist_ok=True)
plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": 300})
dirs = {0: "drlv3_20261012_2324", 1: "drlv3_s1_20261016_2158", 2: "drlv3_s2_20261017_0459"}   # runs of the paper
for s_, d_ in list(dirs.items()):            # after a fresh recomputation: latest run of each seed
    if not os.path.isdir(os.path.join(RUNS, d_)):
        c_ = sorted(glob.glob(os.path.join(RUNS, f"drlv3_s{s_}_2*")))
        dirs[s_] = os.path.basename(c_[-1])
names = {"amsr-hard": "AMSR2, hard", "amsr-bayes": "AMSR2, bayes", "fusion-hard": "Fusion, hard", "fusion-bayes": "Fusion, bayes"}
fig, axes = plt.subplots(1, 4, figsize=(11, 2.9))
for ax, v in zip(axes, names):
    for s, d in dirs.items():
        c = pd.read_csv(os.path.join(RUNS, d, "curves.csv")); c = c[c.variant == v]
        ax.plot(c.steps / 1e6, 100 * c.val_regret, lw=1.1, label=f"seed {s}", color=["#1b9e77", "#d95f02", "#7570b3"][s])
    ax.set_title(names[v], fontsize=9); ax.set_xlabel("PPO steps, million")
axes[0].set_ylabel("validation regret, %"); axes[0].legend(fontsize=7)
fig.tight_layout(); fig.savefig(os.path.join(OUT, "figS1_ppo_curves_en.png"), bbox_inches="tight"); plt.close(fig)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)


def H(t, l=1): doc.add_heading(t, level=l)
def P(t, italic=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.italic = italic; return p
def T(caption, rows):
    P(caption, italic=True)
    tb = doc.add_table(rows=len(rows), cols=len(rows[0])); tb.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tb.cell(i, j); cell.text = str(v)
            for pp in cell.paragraphs:
                for rr in pp.runs:
                    rr.font.size = Pt(9); rr.bold = (i == 0)
    doc.add_paragraph()


doc.add_heading("Supplementary material", 0)
P("Sensor-error-aware ice routing using shipborne imagery and AMSR2: a simulation study based on Polarstern cruise PS131. "
  "Section, table, figure and reference numbers without the prefix S refer to the main text. Differences are given in "
  "percentage points (pp); * marks differences whose 95% confidence interval (episode bootstrap) excludes zero, without "
  "correction for multiple comparisons.", italic=True)

H("S1. Robustness to simulation assumptions")
T("Table S1. Robustness of the results to the simulation assumptions, pp: regret difference fusion − AMSR2 (F−A, bayes "
  "planner) by visibility class, the bayes − hard gain for fusion, and the cvar − bayes difference for AMSR2 / fusion / "
  "camera; the last column gives the change in nominal transit time bayes − hard for fusion, h. Variants with a different "
  "α level change only the cvar planner.",
  [["Variant", "F−A, all", "F−A, > 10 km", "F−A, < 1 km", "bayes−hard (fusion)", "cvar−bayes (A / F / C)", "Time bayes−hard (fusion), h"],
   ["Main run", "−0.75", "−4.7*", "+2.1*", "−13.6*", "−1.8* / −2.3* / −2.1*", "−1.86*"],
   ["Sampling from confusion matrices", "−0.85", "−4.4*", "+1.0*", "−11.4*", "−1.5* / −1.8* / −1.2", "−1.55*"],
   ["Floe clusters σ = 3 km", "−0.13", "−2.7*", "+1.5*", "−14.9*", "−2.2* / −2.5* / −0.4", "−1.88*"],
   ["Floe clusters σ = 12 km", "−0.32", "−3.9*", "+1.8*", "−12.8*", "−2.0* / −2.5* / −0.4", "−1.82*"],
   ["Thresholds 0.25 / 0.75", "−0.34", "−3.2*", "+1.6*", "−14.0*", "−2.4* / −2.7* / −0.7", "−1.94*"],
   ["Thresholds 0.45 / 0.90", "−0.94", "−5.1*", "+1.5*", "−14.2*", "−1.8* / −2.6* / −1.1", "−1.90*"],
   ["Costs 1 / 4 / 8", "−1.46", "−5.7*", "+2.1*", "−7.3*", "−1.6* / −0.9 / −1.8*", "−0.87*"],
   ["Costs 1 / 2 / 12", "−1.00", "−7.5*", "+3.4*", "−21.3*", "−3.1* / −4.8* / −1.7", "−2.77*"],
   ["Costs 1 / 2 / 4 (time-optimal)", "−0.68", "−2.6*", "+0.9*", "−5.2*", "−1.1* / −1.2* / −0.6", "−0.71*"],
   ["Camera range ×2", "−2.11*", "−6.0*", "+0.8*", "−14.4*", "−1.8* / −2.1* / −1.3", "−1.96*"],
   ["Camera error block 1 km", "−3.66*", "−9.5*", "−0.2", "−13.5*", "−2.6* / −2.3* / +1.8*", "−1.92*"],
   ["Camera error block 6 km", "+1.16*", "+1.0", "+1.8*", "−14.1*", "−2.5* / −2.5* / −1.3", "−1.79*"],
   ["Frames from one donor episode", "−0.73", "−4.8*", "+2.1*", "−13.2*", "−1.8* / −2.4* / −1.7*", "−1.82*"],
   ["CVaR, α = 0.5", "—", "—", "—", "—", "−2.3* / −2.7* / −1.0", "—"],
   ["CVaR, α = 0.9", "—", "—", "—", "—", "−2.5* / −2.6* / −1.4", "—"]])
P("Donor episode. In this variant all camera and fusion frames of one trial are drawn from a single randomly chosen episode "
  "with the same visibility class (among episodes outside the group of the evaluated episode). If the donor has no frames "
  "of the required true class, the frame is drawn from a random episode, as in the main observation model. Because the ice "
  "label is nearly constant within an hourly episode, a donor can supply frames of its own class only: 43% of the blocks "
  "were taken from the donor and the remaining 57% from the fallback pool. The random generator used to pick the donor is "
  "separate from the main one, so the true fields, start–goal pairs and AMSR2 observations are identical to the main run "
  "and the differences between the two variants are paired.")
T("Table S2. Correlation of classification errors along the ship track as a function of the distance between frames "
  "(frames with an ice label, out-of-fold predictions). “Pooled” uses all frame pairs within an episode relative to the "
  "sample-wide error rate (and therefore includes the between-episode component); “within episode” uses deviations from "
  "the episode's own error rate.",
  [["Distance, km", "0–0.5", "0.5–1", "1–2", "2–3", "3–5", "5–8", "8–12", "12–20", "20–30"],
   ["Camera, pooled", "+0.72", "+0.52", "+0.51", "+0.53", "+0.49", "+0.46", "+0.54", "+0.49", "+0.41"],
   ["Camera, within episode", "+0.12", "−0.03", "−0.05", "+0.00", "−0.07", "−0.12", "−0.15", "−0.17", "+0.00"],
   ["Fusion, pooled", "+0.51", "+0.36", "+0.36", "+0.43", "+0.33", "+0.23", "+0.26", "+0.14", "+0.23"],
   ["Fusion, within episode", "+0.07", "−0.01", "+0.00", "+0.01", "−0.10", "−0.12", "−0.06", "+0.01", "+0.00"]])
P("The camera error rate is 29.3% of all labelled frames; the median over episodes is 9%, and in 10 of 39 episodes the "
  "camera is wrong on more than half of the frames. For fusion the error rate is 19.0%, the median 4%, and 7 episodes have "
  "an error rate above 50%. The estimate is one-dimensional (along the track) and does not determine the two-dimensional "
  "error correlation along a route.")

H("S2. Statistics of the main run")
T("Table S3. Holm correction for the nine key comparisons of the main run: regret difference (pp), p-value (two-sided "
  "episode bootstrap, 10,000 resamples; smallest resolvable p = 2·10⁻⁴) and Holm-adjusted p.",
  [["Comparison", "Difference", "p", "p (Holm)", "Conclusion"],
   ["AMSR2: bayes − hard", "−6.79", "< 0.001", "< 0.001", "significant"],
   ["Fusion: bayes − hard", "−13.64", "< 0.001", "< 0.001", "significant"],
   ["Camera: bayes − hard", "−20.12", "< 0.001", "< 0.001", "significant"],
   ["AMSR2: cvar − bayes", "−1.83", "< 0.001", "< 0.001", "significant"],
   ["Fusion: cvar − bayes", "−2.33", "< 0.001", "< 0.001", "significant"],
   ["Camera: cvar − bayes", "−2.13", "0.025", "0.068", "not significant"],
   ["Fusion − AMSR2, hard", "+6.09", "< 0.001", "< 0.001", "significant"],
   ["Fusion − AMSR2, bayes", "−0.75", "0.37", "0.37", "not significant"],
   ["Fusion − AMSR2, cvar", "−1.25", "0.023", "0.068", "not significant"]])
rows4 = [["Reference", "Source", "Planner", "Mean, %", "p90, %", "p95, %", "CVaR₀.₉, %", "Heavy ice, km", "Heavy ice p95, km"]]
for ref, data in (("AMSR2 + floes", [
        ["AMSR2", "hard", "29.1", "67.1", "82.9", "88.2", "74.2", "125.8"], ["AMSR2", "bayes", "22.3", "55.7", "72.2", "76.0", "69.4", "114.4"],
        ["AMSR2", "cvar", "20.5", "53.2", "70.3", "76.3", "67.9", "112.0"], ["Fusion", "hard", "35.2", "70.5", "88.9", "97.3", "78.3", "137.3"],
        ["Fusion", "bayes", "21.6", "50.6", "65.3", "76.2", "68.9", "114.9"], ["Fusion", "cvar", "19.2", "49.4", "67.9", "71.9", "67.2", "112.0"],
        ["Camera", "hard", "49.8", "101.7", "145.5", "162.7", "85.0", "134.1"], ["Camera", "bayes", "29.6", "75.4", "110.2", "130.3", "71.4", "112.8"],
        ["Camera", "cvar", "27.5", "70.9", "97.8", "119.2", "70.9", "112.4"]]),
                  ("DMI ≥ 70%", [
        ["AMSR2", "hard", "12.2", "25.1", "32.2", "40.3", "86.2", "133.3"], ["AMSR2", "bayes", "6.7", "14.9", "21.3", "30.3", "81.0", "122.9"],
        ["AMSR2", "cvar", "4.2", "12.9", "20.8", "25.5", "78.9", "119.9"], ["Fusion", "hard", "20.0", "34.6", "42.0", "51.2", "92.7", "143.0"],
        ["Fusion", "bayes", "7.6", "17.6", "23.3", "32.6", "81.7", "123.8"], ["Fusion", "cvar", "4.6", "14.7", "20.9", "26.6", "79.2", "120.3"],
        ["Camera", "hard", "24.7", "40.0", "54.5", "61.4", "94.8", "139.9"], ["Camera", "bayes", "4.8", "15.4", "26.0", "31.5", "79.3", "119.4"],
        ["Camera", "cvar", "4.5", "14.4", "23.8", "31.7", "79.2", "119.0"]])):
    rows4 += [[ref, *r] for r in data]
T("Table S4. Tail metrics of the routes (over all trials): mean regret, 90th and 95th percentiles, mean over the worst 10% "
  "of routes (CVaR₀.₉), and mean and 95th percentile of the heavy-ice distance along the route.", rows4)

H("S3. Alternative reference: DMI ice charts")
T("Table S5. Sensitivity of the DMI-reference results to the treatment of no-data cells and to chart time lag: regret "
  "differences, pp (31 episodes with window coverage ≥ 70% unless stated otherwise). A — AMSR2, F — fusion, C — camera.",
  [["Variant", "bayes − hard (A / F / C)", "cvar − bayes (A / F / C)", "F − A, bayes", "C − A, bayes", "Camera regret, bayes, %"],
   ["No data — impassable (main)", "−5.5* / −12.4* / −19.9*", "−2.5* / −3.0* / −0.4", "+0.85*", "−1.9*", "4.8"],
   ["No data — open water", "−5.6* / −12.6* / −20.8*", "−2.3* / −2.6* / −0.2", "+0.81*", "+3.7", "12.2"],
   ["No data — heavy ice", "−5.3* / −12.4* / −19.7*", "−2.5* / −2.9* / −0.1", "+0.68*", "−0.6", "6.0"],
   ["No data — from AMSR2", "−6.0* / −13.2* / −19.3*", "−2.0* / −2.3* / −1.2*", "+0.39*", "+2.0", "10.2"],
   ["From AMSR2, coverage ≥ 50% (35 ep.)", "−4.9* / −11.7* / −17.0*", "−1.8* / −1.5* / −1.5*", "−0.05", "+2.0", "10.4"],
   ["Chart lag ≤ 24 h (22 ep.)", "−5.1* / −12.3* / −20.1*", "−2.4* / −2.7* / −0.3", "+0.75*", "−1.8*", "—"]])
P("When no-data cells are filled with open water, heavy ice or AMSR2-based classes, 2–22% of the window is filled. The "
  "AMSR2 field is aligned with the chart window using the orientation that gives the highest correlation with the chart. "
  "In all variants 620 start–goal pairs were simulated (700 for coverage ≥ 50%) and all routes reached the goal; since "
  "impassable versus filled no-data cells change the grid connectivity, planner differences rather than absolute regret "
  "should be compared across variants.")

H("S4. Training of the PPO agent")
P("Figure S1. Validation regret (150 scenarios from training episodes) during PPO training for three training seeds. The "
  "episode split and the test scenarios are identical for all seeds; the model with the best validation regret is "
  "evaluated on the 13 held-out episodes (Table 8 of the main text). For no seed does the validation regret decrease "
  "consistently; in many cases it increases as training proceeds, and the best validation model remains close to the "
  "initial policy.", italic=True)
doc.add_picture(os.path.join(OUT, "figS1_ppo_curves_en.png"), width=Cm(16.5))
doc.add_paragraph()

H("S5. Mapping of WMO codes to the classes used in this study")
P("Codes are taken from the synoptic reports in FM 13 SHIP format [47] and decoded according to WMO-No. 306 [48].")
T("Table S6. Horizontal visibility VV (code table 4377), maritime codes 90–99. Codes 90 and 94 do not occur in the data.",
  [["Code", "Visibility", "Class in this study"], ["90", "< 0.05 km", "< 1 km"], ["91", "0.05 km", "< 1 km"], ["92", "0.2 km", "< 1 km"],
   ["93", "0.5 km", "< 1 km"], ["94", "1 km", "1–10 km"], ["95", "2 km", "1–10 km"], ["96", "4 km", "1–10 km"],
   ["97", "10 km", "≥ 10 km"], ["98", "20 km", "≥ 10 km"], ["99", "≥ 50 km", "≥ 10 km"]])
T("Table S7. Concentration or arrangement of sea ice ci (code table 0639).",
  [["Code", "WMO meaning", "Class in this study"],
   ["0", "No sea ice in sight", "≤ 3/10"],
   ["1", "Ship in an open lead more than 1 n mile wide, or ship in fast ice with boundary beyond limit of visibility", "≤ 3/10 (absent in data)"],
   ["2", "Sea ice present in concentrations less than 3/10, open water or very open pack ice", "≤ 3/10"],
   ["3", "4/10 to 6/10, open pack ice", "4–8/10"], ["4", "7/10 to 8/10, close pack ice", "4–8/10"],
   ["5", "9/10 or more, but not 10/10, very close pack ice", "≥ 9/10"],
   ["6", "Strips and patches of pack ice with open water between", "≥ 9/10 (absent in data)"],
   ["7", "Strips and patches of close or very close pack ice with areas of lesser concentration between", "≥ 9/10 (ambiguous, Sect. 3.2)"],
   ["8", "Fast ice with open water, very open or open pack ice to seaward of the ice boundary", "≥ 9/10 (absent in data)"],
   ["9", "Fast ice with close or very close pack ice to seaward of the ice boundary", "≥ 9/10 (absent in data)"],
   ["/", "Unable to report", "no label"]])
P("The present-weather code ww (code table 4677) was used only to describe the imaging conditions and does not enter the model.")
doc.save(os.path.join(OUT, "supplementary_material_en.docx"))
print("ok")
